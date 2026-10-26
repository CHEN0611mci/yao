# -*- coding: utf-8 -*-
"""Build the SRS DOCX from the constrained Markdown source.

Supported markdown subset: cover block, TOC marker, headings 1-4, paragraphs,
bullet lists, numbered lists, pipe tables, fenced code blocks, images, blockquotes
and explicit page-break markers.
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOC_DIR = os.environ.get("SRS_DOC_DIR", os.path.join(BASE, "docs", "01-需求文档"))
DOC_NO = os.environ.get("SRS_DOC_NO", "PET-MALL-SRS-001")
DOC_VER = os.environ.get("SRS_DOC_VER", "V1.0")
MD = os.environ.get("SRS_MD", os.path.join(DOC_DIR, "宠物商城系统_企业级需求文档_v1.0.md"))
OUT = os.environ.get("SRS_OUT", os.path.join(DOC_DIR, "宠物商城系统_企业级需求文档_v1.0.docx"))

NAVY = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x66, 0x66, 0x66)
HEAD_BG = "1F4E79"
ZEBRA_BG = "F2F6FA"
CODE_BG = "F4F4F4"
BODY_W = Cm(15.4)


# ----------------------------------------------------------------- helpers
def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_table_borders(table, color="B7C6D8", size=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def run_font(run, east="宋体", latin="Times New Roman", size=12, bold=False,
             color=DARK, italic=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), east)


def para_shading(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def para_border_left(paragraph, color="1F4E79", size=18):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    pPr.append(borders)


def add_field(paragraph, instr, placeholder="", dirty=False):
    """Insert a Word field as a correct run sequence; returns all created runs."""
    runs = []
    r = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        fld_begin.set(qn("w:dirty"), "true")
    r._r.append(fld_begin)
    runs.append(r)

    r = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    r._r.append(instr_el)
    runs.append(r)

    r = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    r._r.append(fld_sep)
    runs.append(r)

    r = paragraph.add_run(placeholder)
    runs.append(r)

    r = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._r.append(fld_end)
    runs.append(r)
    return runs


def enable_update_fields(document):
    settings = document.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


# ------------------------------------------------------------ inline parse
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=12, bold=False, color=DARK,
               east="宋体", latin="Times New Roman"):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run_font(run, east=east, latin=latin, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run_font(run, east="宋体", latin="Consolas", size=size - 1.5,
                     bold=bold, color=RGBColor(0x9C, 0x27, 0x00))
        else:
            run = paragraph.add_run(part)
            run_font(run, east=east, latin=latin, size=size, bold=bold, color=color)


# ---------------------------------------------------------------- document
def build_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")

    specs = {
        "Heading 1": (16, RGBColor(0x1F, 0x4E, 0x79), 18, 10),
        "Heading 2": (14, RGBColor(0x1F, 0x4E, 0x79), 14, 8),
        "Heading 3": (12.5, RGBColor(0x2E, 0x5E, 0x8C), 10, 6),
        "Heading 4": (12, RGBColor(0x33, 0x33, 0x33), 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.font.name = "Arial"
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), "黑体")
        rFonts.set(qn("w:ascii"), "Arial")
        rFonts.set(qn("w:hAnsi"), "Arial")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def setup_section(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.4)
    section.footer_distance = Cm(1.4)
    section.different_first_page_header_footer = True

    # header: use Normal style so inherited center/right tabs do not interfere
    hp = section.header.paragraphs[0]
    hp.style = doc.styles["Normal"]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.line_spacing = 1.0
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(15.8), alignment=2)
    r1 = hp.add_run(os.environ.get(
        "SRS_HEADER", "萌宠优选 · 宠物商城系统　|　企业级需求文档"))
    run_font(r1, east="宋体", size=9, color=GRAY)
    r2 = hp.add_run("\t" + DOC_NO + "　" + DOC_VER)
    run_font(r2, east="宋体", size=9, color=GRAY)
    pPr = hp._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "B7C6D8")
    borders.append(bottom)
    pPr.append(borders)

    # footer: 第 X 页 / 共 Y 页
    fp = section.footer.paragraphs[0]
    fp.style = doc.styles["Normal"]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(2)
    r = fp.add_run("第 ")
    runs_page = add_field(fp, "PAGE", "1")
    r = fp.add_run(" 页　共 ")
    runs_total = add_field(fp, "NUMPAGES", "1")
    r = fp.add_run(" 页")
    for run in fp.runs:
        run_font(run, east="宋体", size=9, color=GRAY)


def cover(doc, lines, tables):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(lines[0] if lines else "")
    run_font(r, east="黑体", latin="Arial", size=26, bold=True, color=NAVY)

    subtitle = lines[1] if len(lines) > 1 else ""
    main_sub, eng_sub = subtitle, ""
    if "（" in subtitle:
        main_sub, eng_sub = subtitle.split("（", 1)
        eng_sub = eng_sub.rstrip("）")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(main_sub)
    run_font(r, east="黑体", latin="Arial", size=17, bold=True, color=DARK)
    if eng_sub:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(6)
        r = p2.add_run(eng_sub)
        run_font(r, east="宋体", latin="Arial", size=12, bold=False, color=GRAY)

    meta = lines[2] if len(lines) > 2 else ""
    parts = [s.strip() for s in meta.split("｜") if s.strip()]
    if not parts:
        parts = [meta]
    if len(parts) > 2:
        chunks = ["　｜　".join(parts[:2]), "　｜　".join(parts[2:])]
    else:
        chunks = ["　｜　".join(parts)]
    for chunk in chunks:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(chunk)
        run_font(r, east="宋体", size=9.5, color=GRAY)

    # divider
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    borders.append(bottom)
    pPr.append(borders)

    # doc info table (2 columns, no header row)
    rows = tables[0] if tables else []
    rows = [r for r in rows if not (len(r) == 2 and r[0] in ("项目", "字段"))]
    if rows:
        t = doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        t.columns[0].width = Cm(3.8)
        t.columns[1].width = Cm(11.2)
        for i, row in enumerate(rows):
            k, v = row[0], row[1] if len(row) > 1 else ""
            cells = t.add_row().cells
            cells[0].width = Cm(3.8)
            cells[1].width = Cm(11.2)
            pk = cells[0].paragraphs[0]
            pk.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            pk.paragraph_format.space_before = Pt(0.5)
            pk.paragraph_format.space_after = Pt(0.5)
            add_inline(pk, k, size=9.5, bold=True, color=NAVY)
            pv = cells[1].paragraphs[0]
            pv.paragraph_format.space_before = Pt(0.5)
            pv.paragraph_format.space_after = Pt(0.5)
            add_inline(pv, v, size=9.5)
            if i % 2 == 0:
                set_cell_shading(cells[0], ZEBRA_BG)
        set_cell_margins(t, 10, 10, 60, 60)
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "D6E0EC")
            borders.append(el)
        tblPr.append(borders)

    # revision history
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("文档修订历史")
    run_font(r, east="黑体", latin="Arial", size=11, bold=True, color=NAVY)
    if len(tables) > 1:
        add_table(doc, tables[1], widths=None, font_size=8)

    # page break to TOC
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("目　录")
    run_font(r, east="黑体", latin="Arial", size=20, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    runs = add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "（目录域：在 Word 中按 F9 或右键「更新域」生成）", dirty=True)
    for run in runs:
        run_font(run, east="宋体", size=10.5, color=GRAY)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("（本文所有标题均已应用 Word 内置标题样式，更新域后目录将自动生成并带页码）")
    run_font(r, east="宋体", size=9, color=GRAY)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def heading(doc, level, text):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    sizes = {1: 16, 2: 14, 3: 12.5, 4: 12}
    colors = {1: RGBColor(0x1F, 0x4E, 0x79), 2: RGBColor(0x1F, 0x4E, 0x79),
              3: RGBColor(0x2E, 0x5E, 0x8C), 4: RGBColor(0x33, 0x33, 0x33)}
    run_font(r, east="黑体", latin="Arial", size=sizes[level], bold=True,
             color=colors[level])
    if level == 1:
        # thin rule under chapter heading
        pPr = h._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "1F4E79")
        borders.append(bottom)
        pPr.append(borders)


def display_len(s):
    n = 0
    for ch in s:
        n += 2 if ord(ch) > 127 else 1
    return n


def add_table(doc, rows, widths=None, font_size=9):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=0, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    weights = []
    for c in range(ncols):
        w = max(display_len(r[c]) for r in rows)
        weights.append(min(max(w, 10), 44))
    total = sum(weights)
    body_emu = int(BODY_W)
    min_emu = int(Cm(1.9))
    col_widths = [body_emu * w / total for w in weights]
    for _ in range(3):
        deficit = sum(max(0.0, min_emu - x) for x in col_widths)
        if deficit <= 0:
            break
        free_sum = sum(x for x in col_widths if x >= min_emu) or 1.0
        col_widths = [min_emu if x < min_emu else x - deficit * (x / free_sum)
                      for x in col_widths]
    col_widths = [Emu(int(x)) for x in col_widths]

    table.autofit = False
    for ci in range(ncols):
        table.columns[ci].width = col_widths[ci]
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for c, val in enumerate(row):
            cell = cells[c]
            cell.width = col_widths[c]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, val, size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                set_cell_shading(cell, HEAD_BG)
            else:
                add_inline(p, val, size=font_size)
                if i % 2 == 0:
                    set_cell_shading(cell, ZEBRA_BG)
        set_cant_split(table.rows[-1])
    set_repeat_header(table.rows[0])
    set_table_borders(table)
    set_cell_margins(table)

    # spacing paragraph after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.line_spacing = 1.0
    r = sp.add_run("")
    run_font(r, size=4)


def code_block(doc, lines, lang):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.left_indent = Cm(0.4)
    para_shading(p, CODE_BG)
    para_border_left(p, color="8EA9C1", size=12)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line if line else " ")
        run_font(r, east="宋体", latin="Consolas", size=9,
                 color=RGBColor(0x1F, 0x1F, 0x1F))


def image(doc, path, caption):
    if not os.path.exists(path):
        print("WARN missing image:", path)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(path, width=Cm(15.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    run_font(r, east="黑体", latin="Arial", size=10, bold=True,
             color=RGBColor(0x33, 0x33, 0x33))


def body_paragraph(doc, text, first_indent=True):
    p = doc.add_paragraph()
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(24)
    add_inline(p, text)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74 + 0.6 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("• " if level == 0 else "– ")
    run_font(r, size=12, color=NAVY)
    add_inline(p, text)


def numbered(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.95)
    p.paragraph_format.first_line_indent = Cm(-0.95)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.95))
    r = p.add_run(num.rstrip())
    run_font(r, size=12, bold=True, color=NAVY)
    r = p.add_run("\t")
    run_font(r, size=12)
    add_inline(p, text)


def quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    para_border_left(p, color="BF9000", size=18)
    para_shading(p, "FFFBEF")
    add_inline(p, text, size=10.5, color=RGBColor(0x59, 0x50, 0x36))


# ------------------------------------------------------------------ parser
def split_table_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    cover_m, body = None, text
    m = re.search(r"<!-- COVER -->(.*?)<!-- /COVER -->", text, re.S)
    if m:
        cover_m = m.group(1)
        body = text[m.end():]
    toc_marker = "<!-- TOC -->" in body

    elements = []  # (kind, payload)
    lines = body.split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line == "<!-- TOC -->":
            elements.append(("toc", None))
            i += 1
            while i < len(lines) and lines[i].strip() != "<!-- /TOC -->":
                i += 1
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            elements.append(("pagebreak", None))
            i += 1
            continue
        if line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("```"):
            lang = line[3:].strip()
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                block.append(lines[j].rstrip())
                j += 1
            elements.append(("code", (block, lang)))
            i = j + 1
            continue
        mh = re.match(r"^(#{1,6})\s+(.*)$", line)
        if mh:
            elements.append(("h", (len(mh.group(1)), mh.group(2).strip())))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_table_row(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            elements.append(("table", rows))
            continue
        mi = re.match(r"^!\[(.*?)\]\((.*?)\)", line)
        if mi:
            elements.append(("image", (mi.group(2).strip(), mi.group(1).strip())))
            i += 1
            continue
        if line.startswith("> "):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                s = lines[i].strip()[1:].strip()
                buf.append(s)
                i += 1
            elements.append(("quote", " ".join(buf)))
            continue
        if re.match(r"^[-*]\s+", line):
            elements.append(("bullet", re.sub(r"^[-*]\s+", "", line)))
            i += 1
            continue
        mn = re.match(r"^(\d+[.、]\s*)(.*)$", line)
        if mn:
            elements.append(("num", (mn.group(1), mn.group(2))))
            i += 1
            continue
        # paragraph: gather until blank or special
        buf = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("|", "#", "```", ">", "![", "- ", "* "))
                    or re.match(r"^\d+[.、]\s", nxt) or nxt.startswith("<!--")):
                break
            buf.append(nxt)
            i += 1
        elements.append(("p", " ".join(buf)))
    return cover_m, elements, toc_marker


# -------------------------------------------------------------------- main
def main():
    with open(MD, "r", encoding="utf-8") as f:
        md = f.read()
    cover_src, elements, _ = parse(md)

    doc = Document()
    build_styles(doc)
    setup_section(doc)

    # cover
    cov_lines, cov_tables = [], []
    if cover_src:
        cl = [ln.strip() for ln in cover_src.strip().split("\n")]
        i = 0
        while i < len(cl):
            ln = cl[i]
            if not ln:
                i += 1
                continue
            if ln.startswith("|"):
                rows = []
                while i < len(cl) and cl[i].startswith("|"):
                    cells = split_table_row(cl[i])
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                        rows.append(cells)
                    i += 1
                cov_tables.append(rows)
                continue
            cov_lines.append(ln)
            i += 1
    if len(cov_lines) > 3:
        cov_lines = cov_lines[:3]
    cover(doc, cov_lines, cov_tables)

    toc(doc)

    first_chapter = True
    idx = 0
    while idx < len(elements):
        kind, payload = elements[idx]
        if kind == "toc":
            idx += 1
            continue
        if kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            idx += 1
            continue
        if kind == "h":
            level, txt = payload
            if level == 1:
                if not first_chapter:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                first_chapter = False
                heading(doc, 1, txt)
            else:
                heading(doc, min(level, 4), txt)
        elif kind == "p":
            body_paragraph(doc, payload)
        elif kind == "bullet":
            bullet(doc, payload)
        elif kind == "num":
            num, txt = payload
            numbered(doc, num, txt)
        elif kind == "quote":
            quote(doc, payload)
        elif kind == "table":
            add_table(doc, payload)
        elif kind == "code":
            block, lang = payload
            if lang == "mermaid":
                code_block(doc, block, lang)
            else:
                code_block(doc, block, lang)
        elif kind == "image":
            rel, alt = payload
            image(doc, os.path.join(DOC_DIR, rel.replace("/", os.sep)), alt)
        idx += 1

    enable_update_fields(doc)
    doc.save(OUT)
    print("saved", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    sys.exit(main())
